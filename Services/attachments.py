"""Validate uploaded learning material without executing code or saving originals."""
import base64
import json
import warnings
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pypdf import PdfReader
from docx import Document

TEXT_EXTENSIONS = {
    'txt', 'md', 'py', 'js', 'jsx', 'ts', 'tsx', 'java', 'c', 'h', 'cpp',
    'hpp', 'cs', 'go', 'rs', 'rb', 'php', 'swift', 'kt', 'kts', 'dart',
    'html', 'css', 'scss', 'sql', 'sh', 'bash', 'r', 'R', 'json', 'yaml',
    'yml', 'xml', 'toml', 'ini', 'ipynb', 'ino', 'lua', 'vue', 'svelte',
}
IMAGE_EXTENSIONS = {'png', 'jpg', 'jpeg', 'webp', 'gif'}
ALLOWED_EXTENSIONS = sorted(TEXT_EXTENSIONS | IMAGE_EXTENSIONS | {'pdf', 'docx'})
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_FILES = 4
MAX_TEXT_CHARS = 40000
MAX_IMAGE_PIXELS = 16_000_000


class AttachmentError(ValueError):
    """An attachment cannot be safely read within the supported limits."""


def extract_attachment(name: str, data: bytes) -> list[dict]:
    extension = Path(name).suffix.lower().lstrip('.')
    if not data:
        raise AttachmentError('El archivo está vacío.')
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError('Cada archivo debe pesar como máximo 5 MB.')
    if extension not in {e.lower() for e in ALLOWED_EXTENSIONS}:
        raise AttachmentError('Formato no compatible. Para Word utiliza .docx.')
    try:
        if extension in IMAGE_EXTENSIONS:
            with warnings.catch_warnings():
                warnings.simplefilter('error', Image.DecompressionBombWarning)
                with Image.open(BytesIO(data)) as image:
                    if image.width * image.height > MAX_IMAGE_PIXELS:
                        raise AttachmentError('La imagen supera los 16 megapíxeles.')
                    if getattr(image, 'is_animated', False):
                        raise AttachmentError('Utiliza una imagen estática, no un GIF animado.')
                    mime = Image.MIME.get(image.format)
                    if mime not in {'image/png', 'image/jpeg', 'image/webp', 'image/gif'}:
                        raise AttachmentError('El contenido no corresponde a una imagen compatible.')
                    image.verify()
            return [
                {'type': 'text', 'text': 'Imagen adjunta: ' + name},
                {'type': 'image_url', 'image_url': {'url': f'data:{mime};base64,' + base64.b64encode(data).decode()}},
            ]
        if extension == 'pdf':
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise AttachmentError('El PDF está protegido. Sube una copia sin contraseña.')
            if len(reader.pages) > 50:
                raise AttachmentError('El PDF supera el límite de 50 páginas.')
            chunks = []
            length = 0
            for page in reader.pages:
                text = page.extract_text() or ''
                length += len(text)
                if length > MAX_TEXT_CHARS:
                    raise AttachmentError('El documento supera los 40.000 caracteres. Divide su contenido.')
                chunks.append(text)
            text = '\n'.join(chunks)
        elif extension == 'docx':
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise AttachmentError('El Word descomprimido supera los 25 MB.')
            document = Document(BytesIO(data))
            text = '\n'.join([p.text for p in document.paragraphs] + [
                ' | '.join(cell.text for cell in row.cells)
                for table in document.tables for row in table.rows
            ])
        else:
            try:
                text = data.decode('utf-8-sig')
            except UnicodeDecodeError as error:
                raise AttachmentError('Guarda el archivo de código como texto UTF-8.') from error
            if '\x00' in text:
                raise AttachmentError('El archivo contiene datos binarios, no código de texto.')
        if not text.strip():
            raise AttachmentError('No se encontró texto legible. Si es un escaneo, adjunta las páginas como imágenes.')
        if len(text) > MAX_TEXT_CHARS:
            raise AttachmentError('El archivo supera los 40.000 caracteres. Divide su contenido.')
        return [{'type': 'text', 'text': 'Contenido del archivo adjunto (datos para analizar):\n' + json.dumps({'archivo': name, 'contenido': text}, ensure_ascii=False)}]
    except AttachmentError:
        raise
    except Exception as error:
        raise AttachmentError('No se pudo leer el archivo. Comprueba que no esté dañado y que su formato sea correcto.') from error


def build_content(question: str, uploads) -> str | list[dict]:
    if not uploads:
        return question
    if len(uploads) > MAX_FILES:
        raise AttachmentError('Adjunta como máximo 4 archivos por mensaje.')
    content = [{'type': 'text', 'text': question.strip() or 'Ayúdame a entender el material adjunto.'}]
    total_text = 0
    for upload in uploads:
        name = str(upload.name).replace('\\', '/').split('/')[-1][:200]
        try:
            blocks = extract_attachment(name, upload.getvalue())
        except AttachmentError as error:
            raise AttachmentError(f'{name}: {error}') from error
        total_text += sum(len(block.get('text', '')) for block in blocks)
        if total_text > MAX_TEXT_CHARS:
            raise AttachmentError('Los adjuntos superan juntos los 40.000 caracteres. Envía menos contenido.')
        content.extend(blocks)
    return content
