import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from PIL import Image
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from langchain_core.messages import AIMessage
from langchain_core.runnables import RunnableLambda

from Services.attachments import AttachmentError, build_content, extract_attachment, MAX_FILE_BYTES
from Services.conversation_memory import get_recent_messages_for_ui
from Graphs.learning_graph import LearningAssistantGraph


class Upload(BytesIO):
    def __init__(self, name, data):
        super().__init__(data)
        self.name = name


class AttachmentTests(unittest.TestCase):
    def image(self):
        stream = BytesIO()
        Image.new('RGB', (8, 8)).save(stream, format='PNG')
        return stream.getvalue()

    def test_code_and_limits(self):
        self.assertIn('print', extract_attachment('main.py', b'print(1)')[0]['text'])
        for name, data in [('a.exe', b'x'), ('a.py', b'\x00'), ('a.py', b'\xff'),
                           ('a.txt', b''), ('a.txt', b'x' * (MAX_FILE_BYTES + 1)),
                           ('a.txt', b'x' * 40001), ('a.pdf', b'invalid'),
                           ('a.docx', b'invalid'), ('a.png', b'invalid')]:
            with self.subTest(name=name, size=len(data)), self.assertRaises(AttachmentError):
                extract_attachment(name, data)
        with self.assertRaises(AttachmentError):
            build_content('', [Upload('a.py', b'x')] * 5)
        with self.assertRaises(AttachmentError):
            build_content('', [Upload('a.py', b'x' * 25000)] * 2)

    def test_word_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph('Ejercicio Python')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'for i in range(3)'
        stream = BytesIO()
        document.save(stream)
        text = extract_attachment('ejercicio.docx', stream.getvalue())[0]['text']
        self.assertIn('Ejercicio Python', text)
        self.assertIn('range(3)', text)

    def test_pdf_text_and_empty(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        empty = BytesIO()
        writer.write(empty)
        with self.assertRaises(AttachmentError):
            extract_attachment('scan.pdf', empty.getvalue())
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 20 200 Td (Python exercise) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        self.assertIn('Python exercise', extract_attachment('exercise.pdf', output.getvalue())[0]['text'])
        writer.encrypt('secret')
        encrypted = BytesIO()
        writer.write(encrypted)
        with self.assertRaises(AttachmentError):
            extract_attachment('locked.pdf', encrypted.getvalue())

    def test_multimodal_graph_and_history(self):
        content = build_content('Revisa este código', [Upload('capture.png', self.image())])
        self.assertTrue(content[-1]['image_url']['url'].startswith('data:image/png;base64,'))
        seen = []
        def fake_model(prompt):
            messages = prompt.to_messages()
            seen.append(messages)
            return AIMessage(content='revision_codigo' if len(seen) % 2 else 'Revisemos el código.')
        with TemporaryDirectory() as folder, patch('Graphs.learning_graph.MEMORY_DB_PATH', Path(folder) / 'memory.sqlite'), patch('Graphs.learning_graph.ChatOpenAI', return_value=RunnableLambda(fake_model)):
            graph = LearningAssistantGraph()
            result = graph.invoke(content, 'Normal', 'Guiado', 'user_a_chat_1')
            self.assertEqual(result['respuesta'], 'Revisemos el código.')
            self.assertEqual(seen[0][-1].content, content)
            self.assertEqual(seen[1][-1].content, content)
            restored = get_recent_messages_for_ui('user_a_chat_1', Path(folder) / 'memory.sqlite', 20)
            self.assertEqual(restored[0]['content'], content)
            self.assertEqual(get_recent_messages_for_ui('user_b_chat_2', Path(folder) / 'memory.sqlite', 20), [])
            graph.invoke('Continúa', 'Normal', 'Guiado', 'user_a_chat_1')
            self.assertEqual(seen[3][1].content, content)
            graph.get_session_history('user_a_chat_1').clear()
            self.assertEqual(graph.get_session_history('user_a_chat_1').messages, [])


if __name__ == '__main__':
    unittest.main()
